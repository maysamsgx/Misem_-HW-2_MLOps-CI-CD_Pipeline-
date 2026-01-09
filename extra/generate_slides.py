import collections 
import collections.abc
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor

def create_presentation():
    prs = Presentation()

    # Define a clean, professional slide layout function
    def add_slide(title, content_points):
        slide_layout = prs.slide_layouts[1] # Title and Content
        slide = prs.slides.add_slide(slide_layout)
        
        # Title
        title_shape = slide.shapes.title
        title_shape.text = title
        title_shape.text_frame.paragraphs[0].font.size = Pt(36)
        
        # Content
        body_shape = slide.placeholders[1]
        tf = body_shape.text_frame
        tf.word_wrap = True
        
        for point in content_points:
            p = tf.add_paragraph()
            p.text = point
            p.font.size = Pt(20)
            p.space_after = Pt(10)

    # --- SLIDE 1: Title ---
    slide_layout = prs.slide_layouts[0] # Title Slide
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    subtitle = slide.placeholders[1]
    title.text = "JobMatch: Turning Hiring from a Cost Center\ninto a Strategic Engine"
    subtitle.text = "A Level 2 MLOps Ecosystem for Enterprise Scale\nPresented to Senior Management"

    # --- SLIDE 2: Introduction (The Business Problem) ---
    add_slide("The Business Problem: Why This Matters", [
        "\"Stop hiring like it’s 1990. Start scaling like it’s 2030.\"",
        "",
        "The Hard Truth:",
        "- Hiring is no longer an HR problem — it is a BUSINESS RISK.",
        "- High volumes + Manual processes = Inconsistent decisions & Lost Revenue.",
        "",
        "The Core Question:",
        "\"How do we scale hiring decisions with the same rigor, automation, and governance we apply to finance?\""
    ])

    # --- SLIDE 3: The Strategic Solution ---
    add_slide("The Strategic Solution: System > Model", [
        "JobMatch is not just a model. It is a Production-Grade MLOps Ecosystem.",
        "",
        "Designed for High-Velocity Environments:",
        "- Continuously governed decision system.",
        "- Automated pipelines (Training, Deployment, Monitoring).",
        "- Built-in safeguards against unreliable predictions.",
        "",
        "Business Impact: Lower Risk, Predictable Performance, Faster Cycles."
    ])

    # --- SLIDE 4: Executive Value Pillars ---
    add_slide("Executive Value Pillars", [
        "1. Efficiency (Speed):",
        "   - Thousands of profiles processed in milliseconds.",
        "   - Faster staffing -> Faster revenue realization.",
        "",
        "2. Risk Management (Trust):",
        "   - Zero-trust serving logic & Algorithmic Fallback.",
        "   - Defensible decision-making under audit.",
        "",
        "3. Governance (No Black Boxes):",
        "   - Full audit trails (MLflow) & Data Validation (Great Expectations)."
    ])

    # --- SLIDE 5: Designed for Scale ---
    add_slide("Designed for Scale, Not Demos", [
        "Scalability is non-negotiable.",
        "",
        "Technical Enablers:",
        "- High-Cardinality Intelligence: Embeddings for complex skill data.",
        "- Ensemble Architecture: XGBoost + LightGBM + RF for stability.",
        "- Kubernetes-Ready: Horizontal scaling for peak loads.",
        "",
        "Benefit: The system grows with your hiring needs without exponential cost."
    ])

    # --- SLIDE 6: ROI Logic ---
    add_slide("ROI Logic: Why This Investment Makes Sense", [
        "Comparison: Traditional Manual vs. JobMatch System",
        "",
        "Screening Cost:  High (Manual)  ->  Automated (Low)",
        "Time-to-Hire:    Weeks          ->  Minutes",
        "Decision Risk:   Subjective     ->  Quantified & Monitored",
        "Scalability:     Limited        ->  Elastic",
        "Governance:      Minimal        ->  Built-in",
        "",
        "Verdict: This is not just optimization. It is structural improvement."
    ])

    # --- SLIDE 7: Conclusion ---
    add_slide("Strategic Takeaway", [
        "Transformation:",
        "From manual, risk-prone processes -> To a scalable, governed, intelligent system.",
        "",
        "Value for Leadership:",
        "- Efficiency: Hire faster without headcount increase.",
        "- Risk Reduction: Confidence-based decisions.",
        "- Agility: Adapt to market shifts instantly.",
        "",
        "Final Message:",
        "\"JobMatch does not replace judgment. It augments it with speed, consistency, and governance.\""
    ])

    # --- SLIDE 8: Closing ---
    slide_layout = prs.slide_layouts[1]
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    title.text = "Q&A"
    body = slide.placeholders[1]
    tf = body.text_frame
    p = tf.add_paragraph()
    p.text = "Thank You"
    p.font.bold = True
    p.font.size = Pt(32)

    output_path = "Business_Presentation_JobMatch_Executive.pptx"
    prs.save(output_path)
    print(f"Presentation saved to {output_path}")

if __name__ == "__main__":
    create_presentation()
